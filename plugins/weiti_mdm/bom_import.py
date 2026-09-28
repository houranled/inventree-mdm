"""BOM 一键导入核心逻辑：解析文件 → 去重(名称+IPN特征) → 增量建零件 → 挂 BOM。

被插件的 UrlsMixin 视图调用。所有对 InvenTree 模型/插件函数的引用都在函数内
延迟导入，避免模块加载期的循环依赖与 AppRegistryNotReady。
"""

import io
import logging
import os
import re

logger = logging.getLogger('inventree')


# ------------------------------------------------------------------
# 文件解析
# ------------------------------------------------------------------

def _sheet_table(ws):
    """读一个 worksheet → (headers, rows)。空行跳过。"""
    it = ws.iter_rows(values_only=True)
    headers = [str(c).strip() if c is not None else '' for c in next(it, [])]
    rows = []
    for r in it:
        if all(c is None or str(c).strip() == '' for c in r):
            continue
        rows.append({headers[i]: r[i] if i < len(r) else None
                     for i in range(len(headers))})
    return headers, rows


def _sheet_images(ws, zf):
    """一个 worksheet 的全部图片：浮动锚点图 + WPS 单元格嵌入图。"""
    out = {}
    for img in getattr(ws, '_images', []):
        try:
            row = img.anchor._from.row
            fmt = getattr(img, 'format', None) or 'png'
            data = img._data()
        except Exception:
            continue
        if row >= 1 and data:
            out.setdefault(row, (f'row_{row}.{fmt}', data))
    for k, v in _wps_cell_images(ws, zf).items():
        out.setdefault(k, v)
    return out


def _wps_cell_images(ws, zf):
    """WPS 单元格嵌入图：图片在 xl/cellimages.xml，单元格以
    =DISPIMG("ID_xxx",1) 公式引用。返回 {数据行号(1起): (文件名, 字节)}。
    """
    import xml.etree.ElementTree as ET

    if 'xl/cellimages.xml' not in zf.namelist():
        return {}

    # rels: rId -> xl/media/imageN.xxx
    rels = {}
    for r in ET.fromstring(zf.read('xl/_rels/cellimages.xml.rels')):
        rid, tgt = r.get('Id'), (r.get('Target') or '')
        if rid and tgt:
            rels[rid] = tgt if tgt.startswith('xl/') else 'xl/' + tgt.lstrip('/')

    # cellimages.xml: 每个 pic 的 cNvPr@name=DISPIMG id，blip@r:embed=rId
    id2rid = {}
    for pic in ET.fromstring(zf.read('xl/cellimages.xml')).iter():
        if not pic.tag.endswith('}pic'):
            continue
        name = rid = None
        for el in pic.iter():
            tag = el.tag.rsplit('}', 1)[-1]
            if tag == 'cNvPr':
                name = el.get('name')
            elif tag == 'blip':
                rid = el.get(
                    '{http://schemas.openxmlformats.org/officeDocument/'
                    '2006/relationships}embed')
        if name and rid:
            id2rid[name] = rid

    # 单元格公式 -> 行号（openpyxl cell.row 为1基，表头=1 → 数据行号 = row-1）
    out = {}
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if not (isinstance(v, str) and 'DISPIMG' in v):
                continue
            m = re.search(r'DISPIMG\("([^"]+)"', v)
            if not m:
                continue
            path = rels.get(id2rid.get(m.group(1), ''), '')
            if path and path in zf.namelist():
                fmt = path.rsplit('.', 1)[-1]
                out.setdefault(c.row - 1, (f'row_{c.row - 1}.{fmt}',
                                           zf.read(path)))
    return out


def stash_images(images, dest_dir, prefix=''):
    """把 {行号:(文件名,字节)} 落盘，返回 {行号:文件路径}。
    prefix 用于多表导入时区分不同工作表的图片。"""
    os.makedirs(dest_dir, exist_ok=True)
    paths = {}
    for row_i, (fname, data) in images.items():
        fp = os.path.join(dest_dir, f'{prefix}{row_i}_{fname}')
        with open(fp, 'wb') as fh:
            fh.write(data)
        paths[row_i] = fp
    return paths


def parse_all_sheets(django_file):
    """解析文件的全部工作表 → [{'sheet','headers','rows','images'},...]。

    - xlsx: openpyxl 逐表解析，图片逐表提取（浮动图 + WPS DISPIMG）
    - xls:  tablib Databook 逐表解析（不支持图片）
    - csv:  单表，表名取文件名
    无表头或无数据行的工作表会被跳过。
    """
    raw = django_file.read()
    name = (getattr(django_file, 'name', '') or '').lower()

    if name.endswith('.xlsx'):
        import zipfile
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        out = []
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            for ws in wb.worksheets:
                headers, rows = _sheet_table(ws)
                if not headers or not rows:
                    continue
                out.append({'sheet': ws.title, 'headers': headers,
                            'rows': rows, 'images': _sheet_images(ws, zf)})
        return out

    # xls / csv 走 tablib
    import tablib
    if name.endswith('.csv'):
        ds = tablib.Dataset()
        ds.load(raw.decode('utf-8-sig'), format='csv')
        datasets = [ds]
    else:
        book = tablib.Databook()
        book.load(raw, format='xls')
        datasets = list(book.sheets())
    out = []
    for ds in datasets:
        headers = [str(h).strip() if h is not None else ''
                   for h in (ds.headers or [])]
        rows = [{str(k).strip(): row[k] for k in row} for row in ds.dict]
        if headers and rows:
            out.append({'sheet': getattr(ds, 'title', None) or name,
                        'headers': headers, 'rows': rows, 'images': {}})
    return out


def cell(row, col):
    """安全取单元格值 -> 去空白字符串。"""
    if not col:
        return ''
    v = row.get(col)
    if v is None:
        return ''
    return str(v).strip()


# ------------------------------------------------------------------
# 去重 / 特征计算
# ------------------------------------------------------------------

def compute_feature_from_text(category, spec_text):
    """从规格文本按类别模板算出 IPN 特征段。参数不齐返回 None。"""
    from weiti_mdm import category_templates, match_value, param_code
    tpls = category_templates(category)
    if not tpls:
        return None
    segs = []
    for tpl in tpls:
        val = match_value(tpl, spec_text or '')
        if not val:
            return None
        segs.append(param_code(val, tpl))
    return ''.join(segs)


def find_existing_part(name, category, spec_text):
    """按「名称 + IPN特征」查已有零件。

    - 同名且特征段相同(或都无特征) → 命中复用
    - 同名但特征不同 → 视为不同零件(返回 None → 上层新建)
    """
    from part.models import Part
    from weiti_mdm import feature_segment

    same_name = list(Part.objects.filter(name=name))
    if not same_name:
        return None

    incoming_feat = compute_feature_from_text(category, spec_text) if category else None

    for p in same_name:
        try:
            existing_feat = feature_segment(p)
        except Exception:
            existing_feat = None
        if incoming_feat == existing_feat:
            return p
    # 有同名但特征都不匹配：来料无法算特征时，再按规格文本比对——
    # 同名且规格相同才复用；规格不同视为不同零件(如线缆同名不同长度/接口)
    if incoming_feat is None:
        if not spec_text:
            return same_name[0]
        for p in same_name:
            if (p.description or '').strip() == spec_text.strip():
                return p
    return None


# ------------------------------------------------------------------
# 建零件 / 挂 BOM
# ------------------------------------------------------------------

def part_name_from_source(text):
    """从工作表名/文件名取零件名：去扩展名，截掉 'BOM' 及之后内容。

    '间隙传感器BOM清单'                 -> '间隙传感器'
    '间隙采集整车BOM_20260922_含x(1)'    -> '间隙采集整车'
    """
    s = (text or '').strip()
    s = re.sub(r'\.(xlsx|xls|csv)$', '', s, flags=re.I)
    s = re.split(r'BOM', s, flags=re.I)[0]
    return s.strip(' -_—–（）()')


def finished_goods_category():
    """定位「成品」小类：优先 702- 前缀，其次名称含「成品」的编码类别。"""
    from part.models import PartCategory
    from weiti_mdm import CATEGORY_CODE_RE
    cats = [c for c in PartCategory.objects.all().only('name')
            if '成品' in (c.name or '')]
    for c in cats:
        m = CATEGORY_CODE_RE.match(c.name or '')
        if m and m.group(1) == '702':
            return c
    return cats[0] if cats else None


def auto_parent_part(sheet_name):
    """父零件留空时按表名定位/自动创建成品类装配体。

    返回 (part_or_None, action)。action: reused/created/empty。
    """
    from part.models import Part
    name = part_name_from_source(sheet_name)
    if not name:
        return None, 'empty'
    existing = Part.objects.filter(name=name).first()
    if existing:
        if not existing.assembly:
            existing.assembly = True      # 要挂 BOM 行，父件必须是装配体
            existing.save(update_fields=['assembly'])
        return existing, 'reused'
    part = Part(
        name=name,
        description=name,
        category=finished_goods_category(),
        assembly=True, salable=True,
        component=False, purchaseable=False)
    # 保存触发插件信号 → 类别内参数模板决定正式码或 ! 待完善码
    part.save()
    part.refresh_from_db()
    return part, 'created'


def resolve_category(category_text):
    """用插件的归类逻辑把类别文本解析成类别对象（可能 None）。"""
    if not category_text:
        return None
    from weiti_mdm import match_category_from_text
    try:
        return match_category_from_text(category_text)
    except Exception:
        return None


def missing_categories(rows, col_cat):
    """收集类别列里匹配不到系统类别的文本（去重、保持出现顺序）。"""
    seen, out = set(), []
    for row in rows:
        raw = cell(row, col_cat)
        if raw and raw not in seen:
            seen.add(raw)
            if resolve_category(raw) is None:
                out.append(raw)
    return out


def attach_image(part, image_path, dry_run):
    """把图片写进 part.image；零件已有图则跳过。返回附注字符串或 ''。"""
    if not part or not image_path or not os.path.isfile(image_path):
        return ''
    if part.image:
        return '已有图片跳过'
    if dry_run:
        return '将附图片'
    from django.core.files.base import ContentFile
    with open(image_path, 'rb') as fh:
        part.image.save(os.path.basename(image_path), ContentFile(fh.read()))
    return '已附图片'


def get_or_create_part(name, spec_text, category_text, dry_run,
                       image_path=None):
    """返回 (part_or_None, action, note)。action: reused/created/would_create。"""
    category = resolve_category(category_text)
    existing = find_existing_part(name, category, spec_text)
    if existing:
        note = f'复用 #{existing.pk} {existing.IPN or ""}'.strip()
        img_note = attach_image(existing, image_path, dry_run)
        return existing, 'reused', note + (f'；{img_note}' if img_note else '')

    if dry_run:
        cat_name = category.name if category else '(未匹配类别)'
        note = f'将新建 [{cat_name}] {name}'
        if image_path and os.path.isfile(image_path):
            note += '（含图片）'
        return None, 'would_create', note

    from part.models import Part
    part = Part(
        name=name,
        description=spec_text or name,
        category=category,
        keywords=category_text or '',
        component=True,
        purchaseable=True,
    )
    # 保存触发插件信号 → 归类 + IPN + 描述反解参数。
    # 规格反解不全时，assign_ipn 会自动生成 !{code}-{槽位段} 的待完善码，
    # 这里无需再手工补 !。
    part.save()
    part.refresh_from_db()
    note = f'新建 #{part.pk} {part.IPN or ""}'.strip()
    img_note = attach_image(part, image_path, dry_run)
    return part, 'created', note + (f'；{img_note}' if img_note else '')


def add_bom_item(parent, sub_part, quantity, reference, note, dry_run):
    """给父零件挂一个 BOM 行。返回 (ok, msg)。"""
    if dry_run:
        return True, f'将挂 {sub_part or "?"} x{quantity}'
    from part.models import BomItem
    try:
        qty = float(quantity) if quantity else 1.0
    except (ValueError, TypeError):
        qty = 1.0
    # 已存在同 (part, sub_part) 的 BOM 行则更新数量，否则新建
    existing = BomItem.objects.filter(part=parent, sub_part=sub_part).first()
    if existing:
        existing.quantity = qty
        if reference:
            existing.reference = reference
        existing.save()
        return True, f'更新BOM行 {sub_part} x{qty}'
    item = BomItem(
        part=parent, sub_part=sub_part, quantity=qty,
        reference=reference or '', note=note or '')
    item.save()
    return True, f'挂BOM行 {sub_part} x{qty}'


# ------------------------------------------------------------------
# 主流程
# ------------------------------------------------------------------

def run_import(parent_part, rows, mapping, dry_run=True, category_map=None,
               images=None):
    """执行 BOM 导入。

    mapping: {'name':列名, 'qty':列名, 'ref':列名, 'category':列名, 'spec':列名}
    category_map: {Excel原类别文本: 目标类别名}，用户在预览页确认的类别决策
    images: {数据行号(1起): 图片文件路径}，随零件写入 part.image
    返回汇总 dict。
    """
    from django.db import transaction

    col_name = mapping.get('name')
    col_qty = mapping.get('qty')
    col_ref = mapping.get('ref')
    col_cat = mapping.get('category')
    # 规格列可为单列名或列名列表（多列按序拼接为规格文本）
    _spec = mapping.get('spec')
    spec_cols = _spec if isinstance(_spec, list) else ([_spec] if _spec else [])

    report = {
        'created': 0, 'reused': 0, 'bom_rows': 0,
        'failed': 0, 'lines': [],
    }

    def process():
        for i, row in enumerate(rows, start=1):
            name = cell(row, col_name)
            if not name:
                continue
            spec = '; '.join(v for v in
                             (cell(row, c) for c in spec_cols) if v)
            cat_raw = cell(row, col_cat)
            cat = (category_map or {}).get(cat_raw, cat_raw)
            qty = cell(row, col_qty) or '1'
            ref = cell(row, col_ref)

            img_path = (images or {}).get(i)
            try:
                part, action, note = get_or_create_part(
                    name, spec, cat, dry_run, image_path=img_path)
                if action in ('created', 'would_create'):
                    report['created'] += 1
                elif action == 'reused':
                    report['reused'] += 1

                bom_ok, bom_msg = True, ''
                if part is not None:
                    bom_ok, bom_msg = add_bom_item(
                        parent_part, part, qty, ref, spec, dry_run)
                    if bom_ok:
                        report['bom_rows'] += 1
                elif dry_run:
                    # 演练下新零件还没建，BOM 行也计入预期
                    report['bom_rows'] += 1
                    bom_msg = f'将挂 {name} x{qty}'

                report['lines'].append({
                    'row': i, 'name': name, 'action': action,
                    'note': note + (f'；{bom_msg}' if bom_msg else ''),
                    'ok': True,
                })
            except Exception as e:
                logger.exception('WeiTiMDM.bom: 行 %s 失败', i)
                report['failed'] += 1
                report['lines'].append({
                    'row': i, 'name': name, 'action': 'error',
                    'note': str(e), 'ok': False,
                })

    if dry_run:
        # 演练：用事务包住并强制回滚，确保绝不落库
        try:
            with transaction.atomic():
                process()
                raise _Rollback()
        except _Rollback:
            pass
    else:
        with transaction.atomic():
            process()

    return report


class _Rollback(Exception):
    """内部用：演练模式强制回滚。"""
    pass
